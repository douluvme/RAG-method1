import pymupdf
import pytest
from docx import Document

from docqa import config
from docqa.db import Store


@pytest.fixture(autouse=True)
def tmp_data(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DATA_DIR", tmp_path)
    monkeypatch.setattr(config, "FILES_DIR", tmp_path / "files")
    monkeypatch.setattr(config, "DB_PATH", tmp_path / "test.sqlite3")
    return tmp_path


@pytest.fixture
def store(tmp_path):
    return Store(tmp_path / "test.sqlite3")


def make_pdf(path, pages: list[str], scanned_page: int | None = None):
    """A PDF with text pages; `scanned_page` (0-based) holds only an image."""
    doc = pymupdf.open()
    for i, text in enumerate(pages):
        page = doc.new_page()
        if i == scanned_page:
            pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 50, 50), False)
            pix.clear_with(200)
            page.insert_image(pymupdf.Rect(72, 72, 300, 300), pixmap=pix)
        else:
            page.insert_textbox(pymupdf.Rect(72, 72, 520, 770), text, fontsize=11)
    doc.save(path)
    return path


def make_docx(path):
    d = Document()
    d.add_paragraph("Project overview memo.")
    d.add_heading("Budget", level=1)
    d.add_paragraph("The total budget for 2027 is 4.2 million dollars.")
    d.add_paragraph("Marketing receives 15 percent of the budget.")
    d.add_heading("Timeline", level=1)
    d.add_paragraph("The project starts in March 2027.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Phase", "Deadline"
    t.cell(1, 0).text, t.cell(1, 1).text = "Design", "June 2027"
    d.save(path)
    return path


@pytest.fixture
def sample_pdf(tmp_path):
    return make_pdf(
        tmp_path / "contract.pdf",
        [
            "Lease Agreement. The tenant shall pay rent of 1,200 dollars per month.",
            "The term of this agreement is three years, starting on 1 January 2027.",
            "Either party may terminate the agreement with sixty days written notice.",
        ],
    )


@pytest.fixture
def sample_docx(tmp_path):
    return make_docx(tmp_path / "plan.docx")


@pytest.fixture
def sample_txt(tmp_path):
    p = tmp_path / "notes.md"
    lines = [f"Filler line {i}." for i in range(1, 45)]
    lines[41] = "회의는 매주 화요일 오전 10시에 열립니다."
    p.write_text("\n".join(lines), encoding="utf-8")
    return p
